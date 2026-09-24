#!/usr/bin/env python3
"""Module optionnel : collecte d'annonces en ligne -> data/annonces/annonces_<ville>.csv."""
import csv
import os
import sys

import config
from immohunter.annonces.scraper.seloger import scrape
from immohunter.annonces.scraper.geo import filter_by_city, city_to_slug


def dedup(results):
    seen, unique = set(), []
    for r in results:
        url = r.get("url", "")
        key = "{}-{}-{}".format(r.get("prix"), r.get("surface_m2"), r.get("nb_pieces"))
        if url and url not in seen:
            seen.add(url); seen.add(key); unique.append(r)
        elif not url and key not in seen:
            seen.add(key); unique.append(r)
    return unique


def export_csv(results, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    keys = ["plateforme", "titre", "prix", "surface_m2", "nb_pieces", "nb_chambres",
            "localisation", "code_postal", "type_bien", "url", "description"]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=keys, delimiter=";", extrasaction="ignore")
        w.writeheader()
        w.writerows(results)


def export_xlsx(results, path):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        os.makedirs(os.path.dirname(path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Annonces"
        keys = ["plateforme", "localisation", "prix", "surface_m2", "nb_pieces",
                "nb_chambres", "type_bien", "url", "description"]
        hf = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
        hfont = Font(color="FFFFFF", bold=True)
        for ci, k in enumerate(keys, 1):
            c = ws.cell(row=1, column=ci, value=k); c.fill = hf; c.font = hfont
        for ri, r in enumerate(results, 2):
            for ci, k in enumerate(keys, 1):
                val = r.get(k)
                if k == "description" and val:
                    val = str(val)[:500]
                c = ws.cell(row=ri, column=ci, value=val)
                if k == "url" and val:
                    c.hyperlink = str(val)
                    c.font = Font(color="0066CC", underline="single")
                if k == "prix" and isinstance(val, (int, float)):
                    c.number_format = '#,##0'
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        for col in ws.columns:
            mx = max(len(str(c.value or "")[:40]) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(mx + 2, 45)
        wb.save(path)
    except ImportError:
        pass


def main():
    city = config.city

    print("""
+==================================================+
|   IMMO-HUNTER — SeLoger Scraper                   |
+==================================================+
  Ville  : {}
  Type   : {}
  Budget : {} - {}
  Pages  : {}
+==================================================+
""".format(city, config.property_type,
           config.price_min or "tout", config.price_max or "tout", config.max_pages))

    results = scrape()
    if not results:
        print("\nAucune annonce.")
        sys.exit(1)

    local = filter_by_city(results, city)
    print("\n{} annonces a {} (sur {} brutes)".format(len(local), city, len(results)))

    pmin = config.price_min or 0
    pmax = config.price_max or 999999999
    filtered = [r for r in local if pmin <= (r.get("prix") or 0) <= pmax]

    unique = dedup(filtered)
    print("{} uniques".format(len(unique)))

    if not unique:
        print("\nAucune annonce apres filtrage.")
        sys.exit(1)

    slug = city_to_slug(city)
    export_csv(unique, "{}/annonces_{}.csv".format(config.OUTPUT_DIR, slug))
    export_xlsx(unique, "{}/annonces_{}.xlsx".format(config.OUTPUT_DIR, slug))

    with_desc = len([r for r in unique if r.get("description")])
    print("""
  {} — {} annonces exportees
  Avec description : {}
  Excel : data/annonces/annonces_{}.xlsx
  CSV   : data/annonces/annonces_{}.csv
""".format(city, len(unique), with_desc, slug, slug))


if __name__ == "__main__":
    main()
